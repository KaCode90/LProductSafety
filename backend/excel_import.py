"""Read source workbooks without saving them; retain cell/row provenance."""
import hashlib
from collections import Counter
from datetime import datetime, date
from pathlib import Path
import warnings
import openpyxl
from sqlalchemy import select
from backend.store import Record, ImportLog, log

ROOT = Path(__file__).resolve().parents[1]
FILES = ['NFV2_FMD.xlsx','NFV2_TRM.xlsx','NFV2_XRF Test Plan_Product Safety Management.xlsx']


def clean(value):
    if isinstance(value, (datetime,date)):
        return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,str):
        return value.strip()
    return value if value is not None else ''


def read_workbook(path):
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        wb = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)
    entries=[]
    issues=[]
    sheets=[]

    def add(module, sheet, row, mapping, extra=None):
        data = {key:clean(sheet.cell(row,col).value) for key,col in mapping.items()}
        data.update(extra or {})
        data.setdefault('status','Pending')
        source = {'file':path.name,'sheet':sheet.title,'row':row,
                  'cells':{c.coordinate:clean(c.value) for c in sheet[row] if c.value is not None},
                  'formulas':{c.coordinate:c.value for c in formulas[sheet.title][row] if c.data_type=='f'}}
        data['_source']=source
        key=hashlib.sha256(f'{path.name}|{sheet.title}|{row}|{module}'.encode()).hexdigest()
        entries.append((module,key,data))

    for sheet in wb:
        populated=sum(any(c.value is not None for c in r) for r in sheet)
        sheets.append({'name':sheet.title,'rows':sheet.max_row,'populated_rows':populated,'columns':sheet.max_column})
        for row in sheet:
            for c in row:
                if c.data_type=='e':
                    issues.append({'type':'Excel error','sheet':sheet.title,'cell':c.coordinate,'value':str(c.value)})
        name=sheet.title.strip()
        if path.name==FILES[0]:
            for r in range(3,sheet.max_row+1):
                if sheet.cell(r,3).value:
                    add('fmd',sheet,r,dict(project=2,material_code=3,material_name=4,supplier=5,cas=6,substance=7,composition=8,test_type=9,report_id=10,lab=11,issue_date=12,expiry_date=13,flag=14,investigation=15,result=16))
        elif path.name==FILES[1]:
            if name=='TRM Form':
                for r in range(3,sheet.max_row+1):
                    if sheet.cell(r,8).value:
                        add('reports',sheet,r,dict(project=2,material_code=3,material_name=4,supplier=5,manufacturer=6,test_type=7,report_id=8,method=9,issue_date=10,expiry_date=11,result=12,lab=13,file_reference=14,notes=15))
            elif name=='OQC':
                for r in range(3,sheet.max_row+1):
                    if sheet.cell(r,4).value:
                        add('oqc-reports',sheet,r,dict(project=2,test_type=3,report_id=4,method=5,issue_date=6,expiry_date=7,result=8,lab=9,file_reference=10,notes=11))
            elif name=='Test Plan':
                for r in (5,6):
                    add('cts-2',sheet,r,dict(project=1,test=2,category=3,frequency=4,requirement=5),{'requirement_id':f'TRM-PLAN-{r}','notes':'Yêu cầu từ workbook; cần QA review trước khi phê duyệt.'})
                for r in range(9,sheet.max_row+1):
                    if isinstance(sheet.cell(r,1).value,(int,float)) and sheet.cell(r,4).value:
                        add('test-plan',sheet,r,dict(project=2,material_code=3,material_name=4,test_type=5,test=6,category=7,sample_size=8,availability=9,cp=10,oct=11,nov=12,dec=13))
        else:
            if name=='Test Plan':
                for r in (4,5,6):
                    add('cts-2',sheet,r,dict(test=2,category=3,npi_frequency=4,mp_frequency=5),{'requirement_id':f'XRF-FREQUENCY-{r}','requirement':f'{sheet.cell(r,2).value}: NPI {sheet.cell(r,4).value}; MP {sheet.cell(r,5).value}'})
                for r in range(9,sheet.max_row+1):
                    if not isinstance(sheet.cell(r,1).value,(int,float)) or not sheet.cell(r,4).value:
                        continue
                    mapping=dict(project=2,config=3,material_code=4,old_code=5,material_name=6,supplier=7,test=8,category=9,notes=10,thickness=11,sample_size=12)
                    months=['2025-10','2025-11','2025-12']+['2026-'+str(i).zfill(2) for i in range(1,13)]
                    extra={}
                    for c,m in enumerate(months,13):
                        value=sheet.cell(r,c).value
                        extra['plan_'+m]='Planned' if value is True or value==1 else 'Not planned' if value is False or value==0 else clean(value)
                    category=str(sheet.cell(r,9).value or '')
                    idx=4 if category=='Direct Material' else 5 if category=='Indirect Material' else 6
                    extra.update(npi_frequency=clean(sheet.cell(idx,4).value),mp_frequency=clean(sheet.cell(idx,5).value))
                    add('xrf-plan',sheet,r,mapping,extra)
            elif name=='Standard':
                for r in range(2,8):
                    for material,col,spec_col in [('Polymers',6,2),('Metals/Ceramic/Glass',7,None),('Composite',8,3),('Packaging',9,4)]:
                        control=clean(sheet.cell(r,col).value)
                        extra={'material_type':material,'control_limit':control if isinstance(control,(int,float)) else '', 'spec_limit':clean(sheet.cell(r,spec_col).value) if spec_col else '', 'rule':control if isinstance(control,str) else ''}
                        add('xrf-standard',sheet,r,dict(element=5),extra)
                        module,key,data=entries[-1]
                        entries[-1]=(module,hashlib.sha256((key+material).encode()).hexdigest(),data)
            elif name in ('IQC result','OQC result','Change Control'):
                if name=='IQC result':
                    module='xrf-iqc'; mapping=dict(material_code=2,material_name=3,supplier=4,test=5,category=6,build=7,material_type=8,arrival_date=9,lot=10,test_date=11,test_month=12,br=14,cl=15,cd=16,pb=17,cr=18,hg=19,notes=20)
                elif name=='OQC result':
                    module='xrf-oqc'; mapping=dict(project=2,material_code=3,material_name=4,test=5,category=6,notes=7,material_type=8,build=9,lot=10,test_date=11,result=12,br=13,cl=14,cd=15,pb=16,cr=17,hg=18)
                else:
                    module='change-control'; mapping=dict(material_code=2,material_name=3,test=4,category=5,notes=6,material_type=7,build=8,lot=9,test_date=10,result=11,br=12,cl=13,cd=14,pb=15,cr=16,hg=17)
                for r in range(4,sheet.max_row+1):
                    if sheet.cell(r,mapping['material_code']).value and sheet.cell(r,mapping['test_date']).value:
                        add(module,sheet,r,mapping)
            # IQC Trend is a pivot/cache; trend is recomputed from raw IQC measurements.
        if name=='BOM':
            supplier_col=10 if path.name==FILES[1] else 9
            for r in range(2,sheet.max_row+1):
                if sheet.cell(r,8).value:
                    add('bom',sheet,r,dict(customer=1,config=2,parent_code=3,project=4,category=5,level=6,material_code=7,material_name=8,supplier=supplier_col,norm=supplier_col+1))
    wb.close(); formulas.close()
    return entries, {'sheets':sheets,'issues':issues,'counts':dict(Counter(x[0] for x in entries))}


def import_sources(db, actor='System', root=ROOT):
    results=[]
    for filename in FILES:
        path=root/filename
        if not path.exists():
            results.append({'file':filename,'state':'missing'})
            continue
        checksum=hashlib.sha256(path.read_bytes()).hexdigest()
        previous=db.scalar(select(ImportLog).where(ImportLog.filename==filename))
        if previous:
            results.append({'file':filename,'state':'already_imported' if previous.checksum==checksum else 'source_changed_review_required','summary':previous.summary})
            continue
        entries,summary=read_workbook(path)
        for module,key,data in entries:
            db.add(Record(module=module,source_key=key,data=data))
        db.add(ImportLog(filename=filename,checksum=checksum,summary=summary))
        log(db,actor,'Nhập Excel: '+filename,after=summary)
        db.flush()
        results.append({'file':filename,'state':'imported','summary':summary})
    # Build material/supplier masters without overwriting any existing user edits.
    records=db.scalars(select(Record).where(Record.archived==False)).all()
    known={(r.data.get('project'), r.data.get('material_code')) for r in records if r.module=='materials'}
    suppliers={r.data.get('supplier','').casefold() for r in records if r.module=='suppliers'}
    grouped={}
    for record in records:
        d=record.data
        code=str(d.get('material_code','')).strip()
        proj=str(d.get('project','')).strip()
        if record.module == 'bom' and code not in ('','-','N/A'):
            grouped.setdefault((proj, code),[]).append(d)
        supplier=str(d.get('supplier','')).strip()
        if supplier and supplier.casefold() not in suppliers:
            db.add(Record(module='suppliers',data={'supplier':supplier,'status':'Pending'}))
            suppliers.add(supplier.casefold())
    for (proj, code),variants in grouped.items():
        if (proj, code) in known:
            continue
        first=next((v for v in variants if v.get('manufacturer')),variants[0])
        data={k:first.get(k,'') for k in ('material_name','supplier','manufacturer','material_type','project','category')}
        data.update(material_code=code,project=proj,status='Active',approval='Pending',_sources=[v.get('_source',{}) for v in variants])
        db.add(Record(module='materials',data=data))
    db.commit()
    return results


def import_project_bom(db):
    path = ROOT / 'Material report for SE-CM project.xlsx'
    if not path.exists():
        return {'status': 'error', 'message': 'Không tìm thấy file Material report for SE-CM project.xlsx trong thư mục dự án.'}
    wb = openpyxl.load_workbook(path, data_only=True)
    if '01. BOM' not in wb.sheetnames:
        return {'status': 'error', 'message': 'Không tìm thấy sheet "01. BOM" trong file.'}
    ws = wb['01. BOM']
    curr_proj = ''
    curr_cat = ''
    curr_name = ''
    items = []
    for r in range(3, ws.max_row+1):
        p = ws.cell(row=r, column=1).value
        c = ws.cell(row=r, column=2).value
        n = ws.cell(row=r, column=3).value
        s = ws.cell(row=r, column=4).value
        code = ws.cell(row=r, column=5).value
        unit = ws.cell(row=r, column=6).value
        bom_qty = ws.cell(row=r, column=7).value
        if p: curr_proj = str(p).strip()
        if c: curr_cat = str(c).strip()
        if n: curr_name = str(n).strip()
        if not code and not s: continue
        code_str = str(code).strip() if code else ''
        if not code_str: continue
        # Bỏ qua các dòng có định mức BOM = 0 hoặc rỗng
        if bom_qty is None or bom_qty == 0 or str(bom_qty).strip() in ('0', '', '-'):
            continue
        
        sup_str = str(s).strip() if s else ''
        if sup_str:
            sup_str = ' '.join(sup_str.replace('-', ' ').upper().split())

        items.append({
            'project': curr_proj,
            'category': curr_cat,
            'material_name': curr_name,
            'supplier': sup_str,
            'material_code': code_str,
            'unit': str(unit).strip() if unit else '',
            'bom_qty': bom_qty
        })

    now_dt = datetime.now()
    valid_bom_keys = set()
    valid_mat_keys = set()

    valid_sup_keys = set()

    existing_boms = {(b.data.get('project', ''), b.data.get('material_code', ''), b.data.get('supplier', '')): b for b in db.query(Record).filter(Record.module=='bom').all()}
    existing_mats = {(m.data.get('project', ''), m.data.get('material_code', '')): m for m in db.query(Record).filter(Record.module=='materials').all()}
    existing_sups = {r.data.get('supplier'): r for r in db.query(Record).filter(Record.module=='suppliers').all() if r.data.get('supplier')}

    b_count, m_count = 0, 0
    for it in items:
        bom_key = (it['project'], it['material_code'], it['supplier'])
        mat_key = (it['project'], it['material_code'])
        valid_bom_keys.add(bom_key)
        valid_mat_keys.add(mat_key)
        if it['supplier']:
            valid_sup_keys.add(it['supplier'])

        bom_data = {
            'project': it['project'],
            'parent_code': it['project'],
            'category': it['category'],
            'material_code': it['material_code'],
            'material_name': it['material_name'],
            'supplier': it['supplier'],
            'unit': it['unit'],
            'norm': str(it['bom_qty']) if it['bom_qty'] is not None else '',
            'status': 'Active'
        }
        if bom_key in existing_boms:
            b_rec = existing_boms[bom_key]
            b_rec.archived = False
            b_rec.data = bom_data
            b_rec.updated_at = now_dt
        else:
            b_rec = Record(
                module='bom',
                data=bom_data,
                version=1,
                archived=False,
                updated_at=now_dt
            )
            db.add(b_rec)
            existing_boms[bom_key] = b_rec
            b_count += 1

        mat_data = {
            'material_code': it['material_code'],
            'material_name': it['material_name'],
            'supplier': it['supplier'],
            'project': it['project'],
            'category': it['category'],
            'status': 'Active',
            'required_tests': 'RoHS, Halogen-Free'
        }
        if mat_key in existing_mats:
            m_rec = existing_mats[mat_key]
            m_rec.archived = False
            merged_data = dict(m_rec.data)
            merged_data.update({
                'material_code': it['material_code'],
                'material_name': it['material_name'],
                'supplier': it['supplier'],
                'project': it['project'],
                'category': it['category'],
            })
            if not merged_data.get('required_tests'):
                merged_data['required_tests'] = 'RoHS, Halogen-Free'
            if not merged_data.get('status'):
                merged_data['status'] = 'Active'
            m_rec.data = merged_data
            m_rec.updated_at = now_dt
        else:
            m_rec = Record(
                module='materials',
                data=mat_data,
                version=1,
                archived=False,
                updated_at=now_dt
            )
            db.add(m_rec)
            existing_mats[mat_key] = m_rec
            m_count += 1

        if it['supplier']:
            if it['supplier'] not in existing_sups:
                s_rec = Record(module='suppliers', data={'supplier': it['supplier']}, version=1, archived=False, updated_at=now_dt)
                db.add(s_rec)
                existing_sups[it['supplier']] = s_rec
            else:
                existing_sups[it['supplier']].archived = False

    # Archive any BOM records not in this valid BOM set
    for b_key, b_rec in existing_boms.items():
        if b_key not in valid_bom_keys and not b_rec.archived:
            b_rec.archived = True
            b_rec.updated_at = now_dt

    # Archive any Material records not in this valid BOM set
    for m_key, m_rec in existing_mats.items():
        if m_key not in valid_mat_keys and not m_rec.archived:
            m_rec.archived = True
            m_rec.updated_at = now_dt

    # Archive any Supplier records not in this valid BOM set
    for s_key, s_rec in existing_sups.items():
        if s_key not in valid_sup_keys and not s_rec.archived:
            s_rec.archived = True
            s_rec.updated_at = now_dt

    db.commit()
    return {'status': 'ok', 'bom_count': len(valid_bom_keys), 'materials_count': len(valid_mat_keys), 'total_items': len(items)}

